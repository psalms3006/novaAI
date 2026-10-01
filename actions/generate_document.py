"""actions.generate_document — turn content into an actual file on disk.

"Create a PDF report about my project" should end with a PDF the user can
open, not with prose they are invited to paste into Word. So this writes the
file, confirms it exists, and reports where it went and how big it is. If it
cannot, it says so — a generator that reports success for a file nobody can
find is worse than one that refuses.

Content comes from NOVA, not from here. This module's whole job is format and
placement: given text (Markdown-ish), produce a real .txt, .md, .docx, .pdf,
.csv, .xlsx or .pptx in a directory the user actually meant.
"""
from __future__ import annotations

import csv
import io
import logging
import re
from pathlib import Path

log = logging.getLogger(__name__)

#: Formats this module can actually produce. Anything else is refused by name
#: rather than written as a mislabelled text file — a .pdf that is secretly
#: plain text is a lie the user only discovers when they open it.
SUPPORTED = ("txt", "md", "docx", "pdf", "csv", "xlsx", "pptx")

#: Where a document goes when the user did not say. Documents, because that is
#: what the folder is for; never the current working directory, which for a
#: packaged app is wherever the shortcut happened to point.
DEFAULT_FOLDER = "documents"


def _resolve_destination(path: str | None, title: str, fmt: str) -> Path:
    """Work out the real file path from what the user said."""
    from actions.file_controller import _resolve_path, user_folder

    stem = _safe_stem(title) or "nova-document"
    if not path:
        return user_folder(DEFAULT_FOLDER) / f"{stem}.{fmt}"

    target = _resolve_path(path)
    # "save it to Documents" names a folder; "save it as report.pdf" names a
    # file. A trailing suffix is the only reliable difference.
    if target.suffix:
        return target
    return target / f"{stem}.{fmt}"


def _safe_stem(title: str) -> str:
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", (title or "").strip())
    stem = re.sub(r"\s+", " ", stem).strip(" .")
    return stem[:120]


def _split_lines(content: str) -> list[str]:
    return (content or "").replace("\r\n", "\n").split("\n")


# ── writers ──────────────────────────────────────────────────────────────────

def _write_text(path: Path, content: str, title: str) -> None:
    body = content if not title else f"{title}\n{'=' * len(title)}\n\n{content}"
    path.write_text(body, encoding="utf-8")


def _write_markdown(path: Path, content: str, title: str) -> None:
    body = content if not title else f"# {title}\n\n{content}"
    path.write_text(body, encoding="utf-8")


#: A picture on its own line: ![caption](C:/path/figure.jpg or https://...).
#: Asked for "images if you found images" on 2026-09-30, NOVA could only say it
#: had none -- this writer turned every line into text.
_IMAGE_LINE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<ref>[^)]+)\)$")


def _image_file(ref: str):
    """A local picture file for *ref* (a path, or an http(s) URL fetched now)."""
    ref = (ref or "").strip()
    if re.match(r"^https?://", ref, re.I):
        import tempfile
        from actions.research_report import download_image
        return download_image(ref, Path(tempfile.gettempdir()) / "nova_doc_images")
    p = Path(ref)
    return p if p.is_file() else None


def _write_docx(path: Path, content: str, title: str) -> None:
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    if title:
        doc.add_heading(title, level=0)
    for line in _split_lines(content):
        stripped = line.strip()
        if not stripped:
            continue
        pic = _IMAGE_LINE.match(stripped)
        if pic:
            f = _image_file(pic.group("ref"))
            if f is not None:
                try:
                    doc.add_picture(str(f), width=Inches(6))
                except Exception as e:
                    log.warning("picture %s not added: %s", f, e)
                    f = None
            cap = doc.add_paragraph()
            run = cap.add_run(pic.group("alt") if f is not None
                              else f"[picture unavailable] {pic.group('alt')}")
            run.italic, run.font.size = True, Pt(9)
            continue
        # Markdown headings and bullets carry structure worth keeping; a wall
        # of identical paragraphs is not a document.
        if stripped.startswith("#"):
            level = min(len(stripped) - len(stripped.lstrip("#")), 4)
            doc.add_heading(_plain(stripped.lstrip("# ").strip()), level=max(1, level))
        elif stripped[:2] in ("- ", "* "):
            _md_runs(doc.add_paragraph(style="List Bullet"), stripped[2:].strip())
        elif re.match(r"^\d+[.)]\s", stripped):
            _md_runs(doc.add_paragraph(style="List Number"),
                     re.sub(r"^\d+[.)]\s*", "", stripped))
        else:
            _md_runs(doc.add_paragraph(), stripped)
    doc.save(str(path))


def _write_pdf(path: Path, content: str, title: str) -> None:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import (ListFlowable, ListItem, Paragraph,
                                    SimpleDocTemplate, Spacer)

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=11,
                          leading=16, alignment=TA_LEFT, spaceAfter=6)
    flow = []
    if title:
        flow += [Paragraph(_escape(title), styles["Title"]), Spacer(1, 10)]

    bullets: list = []

    def flush_bullets():
        if bullets:
            flow.append(ListFlowable([ListItem(Paragraph(_inline_pdf(b), body))
                                      for b in bullets], bulletType="bullet"))
            flow.append(Spacer(1, 6))
            bullets.clear()

    caption = ParagraphStyle("caption", parent=body, fontSize=9, leading=12,
                             textColor="#555555", spaceAfter=10)
    for line in _split_lines(content):
        stripped = line.strip()
        if not stripped:
            flush_bullets()
            continue
        pic = _IMAGE_LINE.match(stripped)
        if pic:
            flush_bullets()
            f = _image_file(pic.group("ref"))
            if f is not None:
                try:
                    from PIL import Image as _PIL
                    from reportlab.lib.units import cm
                    from reportlab.platypus import Image as RLImage
                    with _PIL.open(f) as im:
                        w, h = im.size
                    width = min(16 * cm, w)
                    flow.append(RLImage(str(f), width=width, height=width * h / w))
                except Exception as e:
                    log.warning("picture %s not added: %s", f, e)
                    f = None
            alt = pic.group("alt") if f is not None else f"[picture unavailable] {pic.group('alt')}"
            flow.append(Paragraph("<i>" + _escape(alt) + "</i>", caption))
            continue
        if stripped.startswith("#"):
            flush_bullets()
            level = min(len(stripped) - len(stripped.lstrip("#")), 3)
            flow.append(Paragraph(_inline_pdf(stripped.lstrip("# ").strip()),
                                  styles[f"Heading{max(1, level)}"]))
        elif stripped[:2] in ("- ", "* "):
            bullets.append(stripped[2:].strip())
        else:
            flush_bullets()
            flow.append(Paragraph(_inline_pdf(stripped), body))
    flush_bullets()

    if not flow:
        flow = [Paragraph("(empty)", body)]
    SimpleDocTemplate(str(path), pagesize=A4,
                      title=title or path.stem).build(flow)


def _escape(text: str) -> str:
    """reportlab reads a mini-markup, so raw & < > would break the build."""
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


#: Inline Markdown the model writes: **bold**, *italic* / _italic_, `code`.
#: It used to reach PDFs as literal asterisks ("**Communicate:**").
_MD_INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|`[^`]+`|\*[^*\s][^*]*\*|(?<![A-Za-z0-9])_[^_\s][^_]*_(?![A-Za-z0-9]))")


def _md_parts(text: str):
    """(kind, text) pieces: kind is '', 'b', 'i' or 'code'."""
    pos = 0
    for m in _MD_INLINE.finditer(text or ""):
        if m.start() > pos:
            yield "", text[pos:m.start()]
        tok = m.group(0)
        if tok[:2] in ("**", "__"):
            yield "b", tok[2:-2]
        elif tok[0] == "`":
            yield "code", tok[1:-1]
        else:
            yield "i", tok[1:-1]
        pos = m.end()
    if pos < len(text or ""):
        yield "", text[pos:]


def _plain(text: str) -> str:
    return "".join(t for _, t in _md_parts(text))


def _inline_pdf(text: str) -> str:
    out = []
    for kind, t in _md_parts(text):
        t = _escape(t)
        out.append({"b": f"<b>{t}</b>", "i": f"<i>{t}</i>",
                    "code": f'<font face="Courier">{t}</font>'}.get(kind, t))
    return "".join(out)


def _md_runs(paragraph, text: str) -> None:
    for kind, t in _md_parts(text):
        run = paragraph.add_run(t)
        if kind == "b":
            run.bold = True
        elif kind == "i":
            run.italic = True
        elif kind == "code":
            run.font.name = "Consolas"


def _rows_from(content: str) -> list[list[str]]:
    """Accept CSV text, or fall back to one column of lines."""
    text = (content or "").strip()
    if not text:
        return [[""]]
    try:
        dialect = csv.Sniffer().sniff(text[:2000])
        return [r for r in csv.reader(io.StringIO(text), dialect)]
    except Exception:
        return [[line] for line in _split_lines(text) if line.strip()]


def _write_csv(path: Path, content: str, title: str) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(_rows_from(content))


def _write_xlsx(path: Path, content: str, title: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = (_safe_stem(title) or "Sheet1")[:31]
    for row in _rows_from(content):
        ws.append(row)
    wb.save(str(path))


def _write_pptx(path: Path, content: str, title: str) -> None:
    from pptx import Presentation

    prs = Presentation()
    if title:
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title

    current = None
    for line in _split_lines(content):
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#") or current is None:
            current = prs.slides.add_slide(prs.slide_layouts[1])
            current.shapes.title.text = _plain(stripped.lstrip("# ").strip()) or "Slide"
        else:
            body = current.placeholders[1].text_frame
            para = body.add_paragraph() if body.text else body.paragraphs[0]
            para.text = _plain(stripped.lstrip("-* ").strip())
    if not prs.slides:
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = title or "Untitled"
    prs.save(str(path))


_WRITERS = {
    "txt": _write_text, "md": _write_markdown, "docx": _write_docx,
    "pdf": _write_pdf, "csv": _write_csv, "xlsx": _write_xlsx,
    "pptx": _write_pptx,
}


# ── entry point ──────────────────────────────────────────────────────────────

def execute(args: dict) -> str:
    content = args.get("content") or ""
    title = (args.get("title") or "").strip()
    fmt = (args.get("format") or "").strip().lower().lstrip(".")
    path_arg = args.get("path") or args.get("destination") or ""

    if not fmt:
        suffix = Path(path_arg).suffix.lower().lstrip(".") if path_arg else ""
        fmt = suffix or "pdf"
    if fmt == "markdown":
        fmt = "md"
    if fmt == "text":
        fmt = "txt"
    if fmt not in SUPPORTED:
        return (f"I can't produce a .{fmt} file. I can make: "
                + ", ".join(SUPPORTED) + ".")
    if not content.strip() and fmt not in ("csv", "xlsx"):
        return "There's no content to put in the document."

    try:
        target = _resolve_destination(path_arg, title, fmt)
    except Exception as e:
        return f"I couldn't work out where to save that: {e}"

    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        _WRITERS[fmt](target, content, title)
    except ImportError as e:
        return (f"I can't write .{fmt} files on this machine — a required "
                f"library is missing ({e}).")
    except Exception as e:
        log.exception("document generation failed")
        return f"I tried to create {target} and it failed: {e}"

    # Say it exists only once it does.
    try:
        if not target.exists():
            return f"FAILED: nothing was written to {target}"
        size = target.stat().st_size
        if size == 0:
            return f"FAILED: {target} was created but is empty"
    except Exception as e:
        return f"FAILED: could not confirm {target} exists ({e})"

    return f"Saved {target} ({size} bytes)."


__all__ = ["execute", "SUPPORTED"]
