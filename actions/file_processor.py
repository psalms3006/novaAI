"""
NOVA file_processor — Handles images, PDFs, documents, CSVs, code files, audio, video, archives.
"""

import os
import json
import traceback
from pathlib import Path
from datetime import datetime


def execute(parameters):
    """
    Main entry point. Parameters dict from Gemini function call.
    """
    file_path = parameters.get("file_path", "")
    action = parameters.get("action", "")
    instruction = parameters.get("instruction", "")
    fmt = parameters.get("format", "")
    width = parameters.get("width")
    height = parameters.get("height")
    scale = parameters.get("scale")
    quality = parameters.get("quality", 85)
    start = parameters.get("start")
    end = parameters.get("end")
    timestamp = parameters.get("timestamp")
    column = parameters.get("column")
    value = parameters.get("value")
    condition = parameters.get("condition", "equals")
    ascending = parameters.get("ascending", True)
    save = parameters.get("save", True)
    destination = parameters.get("destination", "")

    # Resolve file path
    if not file_path:
        return "Error: file_path is required"

    path = Path(file_path)
    if not path.exists():
        return f"Error: File not found: {file_path}"

    ext = path.suffix.lower()
    result = f"Processed {path.name}"

    try:
        # ── Images ───────────────────────────────────────────────────────
        if ext in ('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.webp', '.tiff'):
            from PIL import Image

            if action == "describe" or action == "info":
                img = Image.open(path)
                result = f"Image: {img.size[0]}x{img.size[1]}, mode={img.mode}, format={img.format}"

            elif action == "resize":
                img = Image.open(path)
                new_w = width or int(img.width * (scale or 1))
                new_h = height or int(img.height * (scale or 1))
                resized = img.resize((new_w, new_h), Image.LANCZOS)
                out_path = _output_path(path, "resized", fmt or ext)
                resized.save(out_path, quality=quality)
                result = f"Resized to {new_w}x{new_h}, saved to {out_path}"

            elif action == "compress":
                img = Image.open(path)
                out_path = _output_path(path, "compressed", fmt or ext)
                img.save(out_path, quality=quality, optimize=True)
                orig_size = path.stat().st_size
                new_size = out_path.stat().st_size
                result = f"Compressed: {orig_size/1024:.1f}KB → {new_size/1024:.1f}KB ({out_path})"

            elif action == "convert":
                target = fmt or "png"
                out_path = _output_path(path, "converted", f".{target}")
                img = Image.open(path)
                if img.mode in ('RGBA', 'P') and target.lower() in ('jpg', 'jpeg'):
                    img = img.convert('RGB')
                img.save(out_path)
                result = f"Converted to {target}, saved to {out_path}"

            elif action == "ocr":
                try:
                    import pytesseract
                    img = Image.open(path)
                    text = pytesseract.image_to_string(img)
                    result = f"OCR text:\n{text[:500]}{'...' if len(text) > 500 else ''}"
                except ImportError:
                    result = "Error: pytesseract not installed. Run: pip install pytesseract"

            else:
                result = f"Unknown image action: {action}"

        # ── PDFs ─────────────────────────────────────────────────────────
        elif ext == '.pdf':
            if action == "extract_text":
                try:
                    import fitz  # PyMuPDF
                    doc = fitz.open(path)
                    text = ""
                    for page in doc:
                        text += page.get_text()
                    doc.close()
                    result = f"Extracted text ({len(text)} chars):\n{text[:800]}{'...' if len(text) > 800 else ''}"
                except ImportError:
                    result = "Error: PyMuPDF not installed. Run: pip install PyMuPDF"

            elif action == "summarize":
                try:
                    import fitz
                    doc = fitz.open(path)
                    text = ""
                    for page in doc:
                        text += page.get_text()
                    doc.close()
                    result = f"PDF content for summarization:\n{text[:2000]}{'...' if len(text) > 2000 else ''}"
                except ImportError:
                    result = "Error: PyMuPDF not installed."

            elif action == "to_word":
                try:
                    from docx import Document
                    import fitz
                    doc = fitz.open(path)
                    word_doc = Document()
                    for page in doc:
                        word_doc.add_paragraph(page.get_text())
                    doc.close()
                    out_path = _output_path(path, "converted", ".docx")
                    word_doc.save(out_path)
                    result = f"Converted to Word: {out_path}"
                except ImportError:
                    result = "Error: python-docx or PyMuPDF not installed."

            elif action == "info":
                try:
                    import fitz
                    doc = fitz.open(path)
                    result = f"PDF: {len(doc)} pages, metadata: {doc.metadata}"
                    doc.close()
                except ImportError:
                    result = "Error: PyMuPDF not installed."

            else:
                result = f"Unknown PDF action: {action}"

        # ── Word / Text ──────────────────────────────────────────────────
        elif ext in ('.docx', '.doc', '.txt', '.md', '.rtf'):
            if action == "summarize" or action == "extract_text":
                if ext == '.docx':
                    try:
                        from docx import Document
                        doc = Document(path)
                        text = "\n".join([p.text for p in doc.paragraphs])
                        result = f"Document text:\n{text[:1000]}{'...' if len(text) > 1000 else ''}"
                    except ImportError:
                        result = "Error: python-docx not installed."
                else:
                    text = path.read_text(encoding='utf-8', errors='ignore')
                    result = f"Text content:\n{text[:1000]}{'...' if len(text) > 1000 else ''}"

            elif action == "word_count":
                if ext == '.docx':
                    try:
                        from docx import Document
                        doc = Document(path)
                        text = "\n".join([p.text for p in doc.paragraphs])
                    except ImportError:
                        return "Error: python-docx not installed."
                else:
                    text = path.read_text(encoding='utf-8', errors='ignore')
                words = len(text.split())
                result = f"Word count: {words}"

            elif action == "to_bullet":
                if ext == '.docx':
                    try:
                        from docx import Document
                        doc = Document(path)
                        text = "\n".join([p.text for p in doc.paragraphs])
                    except ImportError:
                        return "Error: python-docx not installed."
                else:
                    text = path.read_text(encoding='utf-8', errors='ignore')
                bullets = "\n".join([f"• {line.strip()}" for line in text.split('\n') if line.strip()])
                result = f"Bullet points:\n{bullets[:1000]}"

            else:
                result = f"Unknown document action: {action}"

        # ── CSV / Excel ────────────────────────────────────────────────────
        elif ext in ('.csv', '.xlsx', '.xls'):
            try:
                import pandas as pd

                if ext == '.csv':
                    df = pd.read_csv(path)
                else:
                    df = pd.read_excel(path)

                if action == "analyze" or action == "info":
                    result = f"Shape: {df.shape}\nColumns: {list(df.columns)}\nDtypes:\n{df.dtypes}\nHead:\n{df.head()}"

                elif action == "stats":
                    result = f"Statistics:\n{df.describe()}"

                elif action == "filter":
                    if not column or value is None:
                        return "Error: column and value required for filter"
                    if condition == "equals":
                        filtered = df[df[column] == value]
                    elif condition == "contains":
                        filtered = df[df[column].astype(str).str.contains(str(value), na=False)]
                    elif condition == "gt":
                        filtered = df[df[column] > float(value)]
                    elif condition == "lt":
                        filtered = df[df[column] < float(value)]
                    else:
                        filtered = df[df[column] == value]
                    result = f"Filtered ({len(filtered)} rows):\n{filtered.head(20).to_string()}"

                elif action == "sort":
                    if not column:
                        return "Error: column required for sort"
                    sorted_df = df.sort_values(by=column, ascending=ascending)
                    result = f"Sorted by {column} ({'asc' if ascending else 'desc'}):\n{sorted_df.head(20).to_string()}"

                elif action == "convert":
                    target = fmt or "csv"
                    out_path = _output_path(path, "converted", f".{target}")
                    if target == "csv":
                        df.to_csv(out_path, index=False)
                    elif target == "xlsx":
                        df.to_excel(out_path, index=False)
                    elif target == "json":
                        df.to_json(out_path, orient='records', indent=2)
                    result = f"Converted to {target}: {out_path}"

                else:
                    result = f"Unknown CSV/Excel action: {action}"

            except ImportError:
                result = "Error: pandas not installed. Run: pip install pandas openpyxl"

        # ── JSON / XML ───────────────────────────────────────────────────
        elif ext in ('.json', '.xml'):
            if action == "validate" or action == "format":
                if ext == '.json':
                    data = json.loads(path.read_text(encoding='utf-8'))
                    result = json.dumps(data, indent=2)[:1000]
                else:
                    result = "XML validation requires lxml. Run: pip install lxml"

            elif action == "analyze":
                if ext == '.json':
                    data = json.loads(path.read_text(encoding='utf-8'))
                    if isinstance(data, dict):
                        result = f"Keys: {list(data.keys())}\nType: {type(data).__name__}\nSize: {len(str(data))} chars"
                    elif isinstance(data, list):
                        result = f"List with {len(data)} items\nFirst item: {str(data[0])[:200] if data else 'empty'}"
                else:
                    result = "XML analysis requires lxml"

            elif action == "to_csv" and ext == '.json':
                try:
                    import pandas as pd
                    data = json.loads(path.read_text(encoding='utf-8'))
                    if isinstance(data, list):
                        df = pd.DataFrame(data)
                        out_path = _output_path(path, "converted", ".csv")
                        df.to_csv(out_path, index=False)
                        result = f"Converted to CSV: {out_path}"
                    else:
                        result = "Error: JSON must be a list of objects to convert to CSV"
                except ImportError:
                    result = "Error: pandas not installed."

            else:
                result = f"Unknown JSON/XML action: {action}"

        # ── Code files ─────────────────────────────────────────────────────
        elif ext in ('.py', '.js', '.ts', '.java', '.cpp', '.c', '.h', '.go', '.rs', '.rb', '.php', '.html', '.css', '.sql'):
            code = path.read_text(encoding='utf-8', errors='ignore')

            if action == "explain":
                result = f"Code ({path.name}, {len(code)} chars):\n```{ext.lstrip('.')}\n{code[:1500]}\n```"

            elif action == "review":
                lines = code.split('\n')
                issues = []
                if ext == '.py':
                    if 'import *' in code:
                        issues.append("Avoid `import *` — use explicit imports")
                    if 'except:' in code and 'except Exception' not in code:
                        issues.append("Bare `except:` found — catch specific exceptions")
                result = f"Review of {path.name}:\n" + "\n".join(issues) if issues else f"No obvious issues found in {path.name}."

            elif action == "run" and ext == '.py':
                import subprocess
                try:
                    output = subprocess.run(
                        ['python', str(path)],
                        capture_output=True, text=True,
                        timeout=parameters.get("timeout", 30)
                    )
                    result = f"Exit code: {output.returncode}\nStdout:\n{output.stdout}\nStderr:\n{output.stderr}"
                except Exception as e:
                    result = f"Run failed: {e}"

            else:
                result = f"Code file loaded ({len(code)} chars). Action: {action}"

        # ── Audio ──────────────────────────────────────────────────────────
        elif ext in ('.mp3', '.wav', '.flac', '.aac', '.ogg', '.m4a', '.wma'):
            if action == "info":
                try:
                    from mutagen.mp3 import MP3
                    from mutagen.wave import WAVE
                    if ext == '.mp3':
                        audio = MP3(path)
                    elif ext == '.wav':
                        audio = WAVE(path)
                    else:
                        audio = None
                    if audio:
                        result = f"Duration: {audio.info.length:.1f}s, Bitrate: {audio.info.bitrate}, Sample rate: {audio.info.sample_rate}"
                    else:
                        result = f"Audio file: {path.stat().st_size/1024:.1f}KB"
                except ImportError:
                    result = "Audio info requires mutagen. Run: pip install mutagen"

            elif action == "transcribe":
                try:
                    import whisper
                    model = whisper.load_model("base")
                    transcription = model.transcribe(str(path))
                    result = f"Transcription:\n{transcription['text'][:1000]}"
                except ImportError:
                    result = "Transcription requires openai-whisper. Run: pip install openai-whisper"

            else:
                result = f"Audio file: {path.stat().st_size/1024:.1f}KB. Action: {action}"

        # ── Video ──────────────────────────────────────────────────────────
        elif ext in ('.mp4', '.avi', '.mkv', '.mov', '.wmv', '.flv', '.webm'):
            try:
                from moviepy.editor import VideoFileClip
                clip = VideoFileClip(str(path))

                if action == "info":
                    result = f"Duration: {clip.duration:.1f}s, Size: {clip.size}, FPS: {clip.fps}"

                elif action == "extract_audio":
                    out_path = _output_path(path, "audio", ".mp3")
                    clip.audio.write_audiofile(out_path)
                    result = f"Audio extracted to {out_path}"

                elif action == "extract_frame" and timestamp:
                    t = _parse_time(timestamp)
                    frame = clip.get_frame(t)
                    from PIL import Image
                    img = Image.fromarray(frame)
                    out_path = _output_path(path, f"frame_{timestamp.replace(':', '-')}", ".png")
                    img.save(out_path)
                    result = f"Frame at {timestamp} saved to {out_path}"

                elif action == "trim" and start and end:
                    t1, t2 = _parse_time(start), _parse_time(end)
                    trimmed = clip.subclip(t1, t2)
                    out_path = _output_path(path, "trimmed", ext)
                    trimmed.write_videofile(out_path, codec='libx264', audio_codec='aac')
                    result = f"Trimmed video saved to {out_path}"

                else:
                    result = f"Video: {clip.duration:.1f}s. Action: {action}"

                clip.close()

            except ImportError:
                result = "Video processing requires moviepy. Run: pip install moviepy"

        # ── Archives ───────────────────────────────────────────────────────
        elif ext in ('.zip', '.rar', '.7z', '.tar', '.gz', '.bz2'):
            import zipfile
            import tarfile

            if action == "list":
                if ext == '.zip':
                    with zipfile.ZipFile(path, 'r') as z:
                        result = "Contents:\n" + "\n".join(z.namelist()[:50])
                elif ext in ('.tar', '.gz', '.bz2'):
                    with tarfile.open(path, 'r') as t:
                        result = "Contents:\n" + "\n".join([m.name for m in t.getmembers()[:50]])
                else:
                    result = "Archive listing requires appropriate library"

            elif action == "extract":
                dest = destination or str(path.parent / path.stem)
                os.makedirs(dest, exist_ok=True)
                if ext == '.zip':
                    with zipfile.ZipFile(path, 'r') as z:
                        z.extractall(dest)
                elif ext in ('.tar', '.gz', '.bz2'):
                    with tarfile.open(path, 'r') as t:
                        t.extractall(dest)
                result = f"Extracted to {dest}"

            else:
                result = f"Unknown archive action: {action}"

        # ── Presentations ──────────────────────────────────────────────────
        elif ext in ('.pptx', '.ppt'):
            if action == "summarize" or action == "extract_text":
                try:
                    from pptx import Presentation
                    prs = Presentation(path)
                    texts = []
                    for slide in prs.slides:
                        for shape in slide.shapes:
                            if hasattr(shape, "text"):
                                texts.append(shape.text)
                    full_text = "\n".join(texts)
                    result = f"Presentation text ({len(full_text)} chars):\n{full_text[:1000]}{'...' if len(full_text) > 1000 else ''}"
                except ImportError:
                    result = "Error: python-pptx not installed. Run: pip install python-pptx"
            else:
                result = f"Unknown presentation action: {action}"

        # ── Fallback ───────────────────────────────────────────────────────
        else:
            result = f"Unknown file type: {ext}. Supported: images, PDFs, docs, CSVs, code, audio, video, archives, presentations."

        # Save result to file if requested
        if save and len(str(result)) > 200:
            out_path = _output_path(path, f"{action}_result", ".txt")
            out_path.write_text(str(result), encoding='utf-8')
            result += f"\n\n[Saved to {out_path}]"

    except Exception as e:
        traceback.print_exc()
        result = f"Error processing {path.name}: {str(e)}"

    print(f"[file_processor] {result[:200]}")
    return result


# ── Helpers ─────────────────────────────────────────────────────────────────

def _output_path(original: Path, suffix: str, ext: str) -> Path:
    """Generate output path like: file_resized.png, file_converted.csv, etc."""
    stem = original.stem
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    name = f"{stem}_{suffix}_{timestamp}{ext}"
    return original.parent / name


def _parse_time(ts: str) -> float:
    """Parse 'HH:MM:SS' or seconds string to float seconds."""
    if ':' in ts:
        parts = ts.split(':')
        if len(parts) == 3:
            h, m, s = map(float, parts)
            return h * 3600 + m * 60 + s
        elif len(parts) == 2:
            m, s = map(float, parts)
            return m * 60 + s
    return float(ts)