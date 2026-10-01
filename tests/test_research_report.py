"""research_report: the five-page, sourced, illustrated report the 2026-09-30
session asked for three times and got as two pages without a picture.

Everything outside the machine is replaced (search, page fetch, the model,
image download); the documents and the pictures in them are real files."""
from __future__ import annotations

import json
import re
from pathlib import Path

import pytest

from actions import generate_document, research_report as rr


def _png(path: Path, w=640, h=360) -> Path:
    from PIL import Image
    Image.new("RGB", (w, h), (30, 90, 160)).save(path, "PNG")
    return path


@pytest.fixture
def offline(tmp_path, monkeypatch):
    calls = {"generate": [], "search": [], "fetch": []}

    def generate(prompt, want_json=False):
        calls["generate"].append(prompt)
        if want_json:
            return json.dumps({
                "title": "Trillion AI and Zoe OS",
                "queries": ["Trillion AI", "Zoe OS", "Zoe OS pricing"],
                "sections": [{"heading": h, "covers": h} for h in
                             ("Overview", "Trillion AI", "Zoe OS", "Comparison", "Analysis and NOVA's assessment")],
                "image_queries": ["Zoe OS screenshot"]})
        m = re.search(r"needs about (\d+)", prompt) or re.search(r"about (\d+) words", prompt)
        n = int(m.group(1)) if m else 300
        return " ".join(["Zoe OS runs agents across devices [1]."] * (n // 7 + 1))

    def search(q, n=6):
        calls["search"].append(q)
        return [{"title": f"{q} page", "url": f"https://example.com/{q.replace(' ', '-')}", "snippet": "s"},
                {"title": "shared", "url": "https://example.com/shared", "snippet": "s"}]

    def fetch(url):
        calls["fetch"].append(url)
        if "pricing" in url:
            raise IOError("HTTP 403")
        return "<p>page text</p>" * 10

    monkeypatch.setattr(rr, "RETRY_PAUSE_S", 0)
    monkeypatch.setattr(rr, "generate", generate)
    monkeypatch.setattr(rr, "search", search)
    monkeypatch.setattr(rr, "fetch_text", fetch)
    monkeypatch.setattr(rr, "image_search", lambda q, n=8: [
        {"title": "Zoe OS [home] screen", "image": "https://img.example/1.png", "page": "https://example.com/zoe"},
        {"title": "tracking pixel", "image": "https://img.example/pixel.gif", "page": "https://x"}])

    def download(url, folder):
        if "pixel" in url:
            return None
        folder.mkdir(parents=True, exist_ok=True)
        return _png(folder / "fig.png")
    monkeypatch.setattr(rr, "download_image", download)
    import actions.file_controller as fc
    monkeypatch.setattr(fc, "user_folder", lambda name: tmp_path / name)
    return calls, tmp_path


def test_report_reaches_the_length_asked_for_with_sources_and_pictures(offline):
    calls, tmp = offline
    rep = rr.build_report("Trillion AI vs Zoe OS", pages=5, workdir=tmp / "work")
    assert rep["words"] >= 5 * rr.WORDS_PER_PAGE * 0.9
    assert rep["sources"] == 4                        # 3 queries + one shared url, deduplicated
    assert rep["figures"] == 1
    md = rep["markdown"]
    assert "## References" in md and "https://example.com/shared" in md
    assert re.search(r"^!\[Figure 1: Zoe OS  home  screen \(source: https://example.com/zoe\)\]\(.+fig\.png\)$",
                     md, re.M)
    assert "Analysis and NOVA's assessment" in md
    assert any("could not be read" in n for n in rep["notes"])


def test_sections_are_written_from_numbered_sources_with_citations(offline):
    calls, tmp = offline
    rr.build_report("Trillion AI vs Zoe OS", pages=2, workdir=tmp / "work")
    section_prompts = [p for p in calls["generate"] if "SOURCES" in p]
    assert section_prompts and all("[1] " in p and "citation" in p for p in section_prompts)
    assert all("never invent" in p for p in section_prompts)


def test_a_short_section_is_expanded_not_shipped(offline, monkeypatch):
    calls, tmp = offline
    real = rr.generate

    def stingy(prompt, want_json=False):
        if "is too short" in prompt or want_json:
            return real(prompt, want_json)
        return "Too short [1]."
    monkeypatch.setattr(rr, "generate", stingy)
    rep = rr.build_report("x", pages=3, workdir=tmp / "w")
    assert rep["words"] >= 3 * rr.WORDS_PER_PAGE * 0.9
    assert sum("is too short" in p for p in calls["generate"]) == 5


def test_execute_saves_a_docx_with_the_picture_in_it(offline):
    out = rr.execute({"topic": "Trillion AI vs Zoe OS", "pages": 5, "format": "docx"})
    assert out.startswith("Saved "), out
    path = Path(out[len("Saved "):].split(": ", 1)[0])
    from docx import Document
    doc = Document(str(path))
    assert len(doc.inline_shapes) == 1
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "References" in text and len(text.split()) >= 2400
    assert "1 picture(s)" in out and "4 sources" in out


def test_execute_saves_a_pdf_and_measures_its_pages(offline):
    out = rr.execute({"topic": "Trillion AI vs Zoe OS", "pages": 5, "format": "pdf"})
    assert out.startswith("Saved "), out
    m = re.search(r"\((\d+) pages\)", out)
    assert m and int(m.group(1)) >= 5


def test_sources_must_name_what_the_report_is_about(offline, monkeypatch):
    """First real run: 'Trillion AI' searches returned trillion-dollar-market
    pages, and every section said nothing could be confirmed."""
    seen = []

    def search(q, n=6):
        seen.append(q)
        if q.startswith('"Trillion AI"'):
            return [{"title": "Trillion AI - voice-first co-founder", "url": "https://trillion.example/", "snippet": ""},
                    {"title": "About Trillion AI", "url": "https://news.example/trillion-ai", "snippet": ""}]
        return [{"title": "AI market to reach $1.3 trillion", "url": "https://market.example/", "snippet": "AI spending"},
                {"title": "The $45 Trillion AI Revolution", "url": "https://sum.example/", "snippet": ""},
                {"title": "380 trillion AI tokens", "url": "https://tokens.example/", "snippet": ""},
                {"title": "Zoe OS launch", "url": "https://zoe.example/", "snippet": "Zoe OS is"},
                {"title": "Zoe OS review", "url": "https://zoe.example/review", "snippet": ""}]
    monkeypatch.setattr(rr, "search", search)
    notes = []
    got = rr._gather([rr._quote_entities("Trillion AI features", ["Trillion AI", "Zoe OS"])],
                     notes, ["Trillion AI", "Zoe OS"])
    assert seen[0] == '"Trillion AI" features'
    urls = [s["url"] for s in got]
    assert not {"https://market.example/", "https://sum.example/", "https://tokens.example/"} & set(urls)
    assert {"https://trillion.example/", "https://zoe.example/"} <= set(urls)
    assert any(q == '"Trillion AI"' for q in seen) or seen[0].startswith('"Trillion AI"')


def test_only_pictures_of_the_subject_are_used(offline, tmp_path, monkeypatch):
    monkeypatch.setattr(rr, "image_search", lambda q, n=8: [
        {"title": "BVITE AI dashboard template", "image": "https://img.example/a.png", "page": "https://dribbble.example/x"},
        {"title": "Zoey OS desktop", "image": "https://img.example/b.png", "page": "https://x.example/"},
        {"title": "screenshot", "image": "https://img.example/c.png", "page": "https://zoeos.example/home"}])
    notes = []
    figs = rr._pictures(["Zoe OS screenshot"], tmp_path / "f", notes, ["Zoe OS", "Zoey OS"])
    assert [f["page"] for f in figs] == ["https://x.example/", "https://zoeos.example/home"]
    figs = rr._pictures(["x"], tmp_path / "g", notes, ["Trillion AI"])
    assert figs == [] and any("no pictures of the subject" in n for n in notes)


def test_a_rate_limited_search_is_retried_once(offline, monkeypatch):
    monkeypatch.setattr(rr, "RETRY_PAUSE_S", 0)
    tries = []

    def flaky(q, n=6):
        tries.append(q)
        if len(tries) == 1:
            raise RuntimeError("ratelimit")
        return [{"title": "Zoe OS", "url": "https://zoe.example/", "snippet": ""}]
    monkeypatch.setattr(rr, "search", flaky)
    notes = []
    assert rr._gather(["Zoe OS"], notes)[0]["url"] == "https://zoe.example/"
    assert len(tries) == 2 and notes == []


def test_searching_stops_at_its_time_budget(offline, monkeypatch):
    monkeypatch.setattr(rr, "SEARCH_BUDGET_S", -1)
    notes = []
    assert rr._gather(["a", "b"], notes) == []
    assert notes == ["searching ran out of time, so some angles were not searched"]


def test_a_report_with_no_sources_fails_honestly(offline, monkeypatch):
    monkeypatch.setattr(rr, "search", lambda q, n=6: [])
    out = rr.execute({"topic": "nothing at all"})
    assert out.startswith("FAILED") and "no web sources" in out


def test_shortfall_is_reported(offline, monkeypatch):
    monkeypatch.setattr(rr, "generate", lambda p, want_json=False: (
        json.dumps({"title": "t", "queries": ["q"], "sections": [{"heading": "a"}], "image_queries": []})
        if want_json else "short [1]."))
    out = rr.execute({"topic": "t", "pages": 5, "images": False})
    assert "shorter than the 5 pages" in out


def test_download_image_refuses_non_http(tmp_path):
    assert rr.download_image("file:///C:/Windows/win.ini", tmp_path) is None
    assert rr.download_image("javascript:alert(1)", tmp_path) is None


def test_generate_document_embeds_local_pictures_and_marks_missing_ones(tmp_path, monkeypatch):
    import actions.file_controller as fc
    monkeypatch.setattr(fc, "user_folder", lambda name: tmp_path)
    pic = _png(tmp_path / "a b.png")                  # a space in the path, as user folders have
    content = (f"# Title\n\nText.\n\n![Figure 1: chart (source: https://x/y)]({pic.as_posix()})\n\n"
               f"![Figure 2: gone]({(tmp_path / 'missing.png').as_posix()})\n")
    out = generate_document.execute({"content": content, "title": "pics", "format": "docx"})
    assert out.startswith("Saved "), out
    from docx import Document
    doc = Document(str(tmp_path / "pics.docx"))
    assert len(doc.inline_shapes) == 1
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Figure 1: chart" in text and "[picture unavailable] Figure 2: gone" in text
    out = generate_document.execute({"content": content, "title": "pics", "format": "pdf"})
    assert out.startswith("Saved "), out


def test_research_report_is_a_declared_permitted_background_tool():
    import nova
    from desk import confirm, live_session
    from nova_core import permissions
    assert any(d["name"] == "research_report" for d in nova.TOOL_DECLARATIONS)
    assert "research_report" in permissions.TOOL_CAPABILITIES
    assert confirm.scope_for("research_report", {}) == "file_write"
    assert "research_report" in live_session.NON_BLOCKING_TOOLS
